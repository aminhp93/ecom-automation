"""Nạp dữ liệu nghiên cứu ban đầu (23–25/09/2026) vào Supabase project "dropship".

Nguồn:
  - pipeline-san-pham-2026-09-23.xlsx (Google Drive) — 91 SP + các sheet dữ liệu thô
  - workspace/doc/1-market-research/7-pipeline-san-pham-2026-09-23/raw/ — JSON/TSV gốc + 2026-09-25.json
  - link-meta-pages.gs — page_id Facebook của từng brand
  - data/eval_2026_09_25.py — quyết định + nhận xét từng SP

Chạy: python3 scripts/research/import_initial.py   (cần hàm tạm _tmp_import trên DB)
"""
import json, os, re, sys, unicodedata, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT.parent / "workspace/doc/1-market-research/7-pipeline-san-pham-2026-09-23/raw"
DRIVE = Path.home() / "Library/CloudStorage/GoogleDrive-minhpham0529@gmail.com/My Drive/Dropship/tai lieu chung"
XLSX = DRIVE / "pipeline-san-pham-2026-09-23.xlsx"
sys.path.insert(0, str(Path(__file__).parent / "data"))
from eval_2026_09_25 import E, CHON_CHINH, CHON_PHU, CHON_DP, KHONG  # noqa: E402

env = dict(l.split("=", 1) for l in (ROOT / ".env.local").read_text().splitlines() if l.startswith("SUPABASE_"))
URL, KEY = env["SUPABASE_URL"], env["SUPABASE_PUBLISHABLE_KEY"]
D23, D24, D25 = "2026-09-23", "2026-09-24", "2026-09-25"


def api(path, body=None):
    req = urllib.request.Request(f"{URL}/rest/v1/{path}", method="POST" if body is not None else "GET",
                                 data=json.dumps(body).encode() if body is not None else None,
                                 headers={"apikey": KEY, "Content-Type": "application/json"})
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read() or "null")


def put(table, rows):
    rows = [{k: v for k, v in r.items()} for r in rows]
    if not rows:
        return
    keys = sorted({k for r in rows for k in r})
    rows = [{k: r.get(k) for k in keys} for r in rows]
    n = 0
    for i in range(0, len(rows), 500):
        n += api("rpc/_tmp_import", {"tbl": table, "rows": rows[i:i + 500]})
    print(f"{table}: +{n}/{len(rows)}")


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).replace("~", "").replace(",", "").replace("%", "").strip()
    try:
        f = float(s)
    except ValueError:
        return None
    return f / 100 if str(v).strip().endswith("%") else f


def to_int(v):
    f = num(v)
    return int(f) if f is not None else None


def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


# ---------- đọc xlsx ----------
wb = load_workbook(XLSX, data_only=True)
ws = wb["Pipeline 91 SP"]
H = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(H)}
pipe = [dict(zip(H, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
assert len(pipe) == 91

DECISION = {CHON_CHINH: "chon_chinh", CHON_PHU: "chon_phu", CHON_DP: "du_phong", KHONG: "khong_chon"}

# ---------- products ----------
products, seen = [], set()
for r in pipe:
    no = int(r["#"])
    slug = slugify(r["Keyword (EN)"])
    while slug in seen:
        slug += "-2"
    seen.add(slug)
    dec = DECISION[E[no][0]]
    products.append(dict(
        slug=slug, name_vi=r["Sản phẩm"], name_en=r["Keyword (EN)"], category=r["Ngành"],
        cluster=r["Cụm store"] or None, stage="deep_dive" if dec != "khong_chon" else "decided",
        decision=dec, decision_summary=E[no][1].split("\n")[0].replace("KẾT LUẬN: ", ""),
        idea_source="Pipeline 91 SP (Topview keyword, 23/09/2026)", legacy_no=no,
        legacy_score=num(r["Điểm (0–100)"]), legacy_status=r["Quyết định"], legacy_reason=r["Lý do"]))

NEW = [  # SP bổ sung từ phân tích 25/09
    dict(slug="teething-roller-snuglet", name_vi="Con lăn / đồ gặm mọc răng (Snuglet)", name_en="teething toys",
         category="Baby", cluster=None, stage="decided", decision="khong_chon",
         decision_summary="Không làm lúc này — ad thắng nhờ claim giảm đau (rủi ro TGA), hàng cho trẻ sơ sinh, giá thấp nếu bỏ claim.",
         idea_source="Store Snuglet (08/2026)", keywords=["teething toys", "teether", "baby teether"]),
    dict(slug="no-pull-dog-harness", name_vi="Dây đai chó chống kéo (no-pull harness)", name_en="no pull dog harness",
         category="Pet", cluster=None, stage="decided", decision="khong_chon",
         decision_summary="Không làm — đối thủ AU có kho Sydney + hoàn tiền 60 ngày, đổi size, CPM cao.",
         idea_source="Anh Thanh đang làm (buổi 24/09)", keywords=["no pull dog harness", "dog harness"]),
    dict(slug="tech-pouch", name_vi="Túi đựng phụ kiện công nghệ (tech pouch)", name_en="tech pouch",
         category="EDC", cluster="B. Store sắp xếp hành lý du lịch", stage="deep_dive", decision="chon_phu",
         decision_summary="SP phụ cho store du lịch (Travel & Everyday Carry Organizer).",
         idea_source="EDC 'sạch' — gợi ý 'EDC gear' của anh Thanh", keywords=["tech pouch"]),
    dict(slug="rfid-wallet", name_vi="Ví chống RFID", name_en="rfid wallet", category="EDC", cluster=None,
         stage="scored", decision=None, decision_summary=None,
         idea_source="EDC 'sạch' — gợi ý 'EDC gear' của anh Thanh", keywords=["rfid wallet", "slim wallet", "mens wallet"]),
]
put("products", [{k: v for k, v in p.items() if k != "keywords"} for p in products + NEW])
ids = {p["slug"]: p["id"] for p in api("products?select=id,slug")}
by_no = {p["legacy_no"]: ids[p["slug"]] for p in products}

# ---------- keywords ----------
kw_rows = [dict(product_id=by_no[int(r["#"])], keyword=r["Keyword (EN)"], is_primary=True) for r in pipe]
for p in NEW:
    kw_rows += [dict(product_id=ids[p["slug"]], keyword=k, is_primary=i == 0) for i, k in enumerate(p["keywords"])]
kw_rows += [dict(product_id=by_no[70], keyword="cable organiser", is_primary=False),
            dict(product_id=by_no[17], keyword="key organiser", is_primary=False)]
put("product_keywords", kw_rows)

# ---------- amazon keyword snapshots ----------
az = []
for r in wb["Amazon keyword (thô)"].iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    mk, kw, se, pu, pr, ap, cc, prod, adp, bid, mo = r[:11]
    az.append(dict(keyword=kw, market=mk, data_month=f"{float(mo or 2026.08):.2f}".replace(".", "-"), captured_on=D23,
                   searches=to_int(se), purchases=to_int(pu), purchase_rate=num(pr), avg_price=num(ap),
                   click_concentration=num(cc), products=to_int(prod), ad_products=to_int(adp), bid=num(bid)))
d25 = json.loads((RAW / "2026-09-25.json").read_text())
for mk, kw, se, pu, pr, ap, cc, prod, adp, bid in d25["amazon_keywords"]:
    az.append(dict(keyword=kw, market=mk, data_month="2026-08", captured_on=D25, searches=se, purchases=pu,
                   purchase_rate=pr, avg_price=ap, click_concentration=cc, products=prod, ad_products=adp, bid=bid,
                   note="purchases = lượt mua từ đúng keyword này, không phải tổng doanh số"))
put("amazon_keyword_snapshots", az)

listings = [dict(zip(d25["amazon_listing_columns"], x), captured_on=D25) for x in d25["amazon_listings"]]
put("amazon_listing_snapshots", listings)

# ---------- trends ----------
series = json.loads((RAW / "trends_us.json").read_text())
metrics = json.loads((RAW / "trends_metrics.json").read_text())
put("trend_snapshots", [dict(keyword=k, market="US", captured_on=D23, median=m.get("median"),
                             floor_ratio=m.get("stability"), seasonality=m.get("seasonality"),
                             growth_25_22=m.get("growth_22_25"), series=series.get(k),
                             note="Monthly 2021-10→2026-10; đỉnh 3–7/2026 là nhiễu hệ thống Topview")
                        for k, m in metrics.items()])

# ---------- tiktok ----------
tt = []
for r in wb["TikTok Shop US 30 ngày (thô)"].iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    kw, tot, top, sold, price, gmv, it2, s2, p2, vurl, plays, note = r[:12]
    tt.append(dict(keyword=kw, market="US", captured_on=D23, items_total=to_int(tot), top_item=top,
                   top_units_30d=to_int(sold), top_price=num(price), top_gmv_30d=num(gmv), item2=it2,
                   item2_units_30d=to_int(s2), item2_price=num(p2), best_video_url=vurl,
                   best_video_views=to_int(plays), note=note))
put("tiktok_snapshots", tt)

# ---------- meta keyword ----------
mt = []
for r in wb["Meta Ads AU (thô)"].iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    kw, st, act, rend, dadv, o60, top, note = r[:8]
    pages = [re.sub(r"\(\d+\)$", "", x.strip()) for x in (top or "").split(";") if x.strip()]
    mt.append(dict(keyword=kw, country="AU", search_type="exact" if st == "exact" else "unordered", captured_on=D23,
                   active_ads=to_int(act), active_ads_approx=str(act or "").startswith("~"), sample_size=to_int(rend),
                   advertisers_in_sample=to_int(dadv), ads_over_60d_in_sample=to_int(o60), top_pages=pages, note=note))
for kw, act, sample, o60, pages in d25["meta_keywords_au_exact"]:
    mt.append(dict(keyword=kw, country="AU", search_type="exact", captured_on=D25, active_ads=act, sample_size=sample,
                   ads_over_60d_in_sample=o60, top_pages=pages, active_ads_approx=False))
put("meta_keyword_snapshots", mt)

# ---------- advertisers ----------
gs = (DRIVE / "link-meta-pages.gs").read_text()
PAGE_IDS = dict(re.findall(r'"([^"]+)":\s*"(\d+)"', gs.split("};")[0]))
KIND = {"advertorial": ["The Better Home Guide", "Jennifer's Travels", "Travel with Alan", "Mia Travel Tips",
                        "Essential Health Finds", "Nerve Support Community", "Gut Reset Project", "Aussie Vitality",
                        "Simplify Living"],
        "marketplace": ["Temu AU", "Temu NZ", "Temu", "Temu Malaysia", "Temu Canada"],
        "retailer": ["BIG W", "Petbarn", "Pillow Talk AU"], "local_service": ["CQ Smart Screens"], "other": ["Macorner"]}
kind_of = {n: k for k, ns in KIND.items() for n in ns}
aj_col = col["Meta pages: active/tổng - năm tạo page"]
adv, snaps, links = {}, [], []
line_re = re.compile(r"^(.+?):\s*([\d,]+)\s*/\s*([\d,]+)\+?\s*-\s*(.+)$")
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or not r[aj_col]:
        continue
    for line in str(r[aj_col]).split("\n"):
        m = line_re.match(line.strip())
        if not m or m.group(1) not in PAGE_IDS:
            continue
        name, a, t, created = m.groups()
        pid = PAGE_IDS[name]
        adv[pid] = dict(page_id=pid, name=name, kind=kind_of.get(name, "brand"),
                        page_created=None if "không" in created else created.strip())
        snaps.append(dict(page_id=pid, captured_on=D24, active_all=int(a.replace(",", "")),
                          total_all=int(t.replace(",", "")), note="Từ cột AJ (tổng tất cả quốc gia)"))
        links.append(dict(product_id=by_no[int(r[0])], page_id=pid, role="competitor"))
for pid, name, kind, act, tot, au, longest, note in d25["advertiser_snapshots"]:
    adv.setdefault(pid, dict(page_id=pid, name=name, kind=kind, page_created=None))
    adv[pid]["kind"] = kind
    if note:
        adv[pid]["note"] = note
    snaps.append(dict(page_id=pid, captured_on=D25, active_all=act, total_all=tot, active_au=au,
                      longest_active_days=longest, note=None))
for a in adv.values():
    a.setdefault("note", None)
put("advertisers", list(adv.values()))
put("advertiser_snapshots", list({(s["page_id"], s["captured_on"]): s for s in snaps}.values()))
EXTRA = {"teething-roller-snuglet": ["672720455915598", "850508544816345"],
         "no-pull-dog-harness": ["107848155017641", "105919255095881", "101251665553495", "102837007868128"],
         "tech-pouch": ["464313140309473"],
         "rfid-wallet": ["113833593676", "381727498962869", "353699631742927", "362145834181902"]}
for slug, pids in EXTRA.items():
    links += [dict(product_id=ids[slug], page_id=p, role="competitor") for p in pids]
links += [dict(product_id=by_no[17], page_id="464313140309473", role="competitor"),
          dict(product_id=by_no[21], page_id="113833593676", role="competitor"),
          dict(product_id=by_no[21], page_id="353699631742927", role="competitor"),
          dict(product_id=by_no[21], page_id="362145834181902", role="competitor")]
put("product_advertisers", list({(l["product_id"], l["page_id"]): l for l in links}.values()))

# ---------- assessments ----------
put("product_assessments", [dict(
    product_id=by_no[int(r["#"])], assessed_on=D23, weight_class=r["Cân nặng [ước tính]"],
    hazmat=r["Pin/lỏng/dễ vỡ/lưỡi dao"], policy_risk=r["Rủi ro chính sách/claim [ước tính]"],
    retail_risk=r["Rủi ro siêu thị Kmart/BigW/Walmart [ước tính]"], bundle_potential=r["Khả năng bundle/AOV [ước tính]"],
    target_dtc_price=r["Giá bán DTC mục tiêu [ước tính]"], landed_cost_note=r["Giá nhập/landed cost"]) for r in pipe] + [
    dict(product_id=ids["teething-roller-snuglet"], assessed_on=D25, weight_class="nhẹ", hazmat="Không",
         policy_risk="Cao", retail_risk="Cao", bundle_potential="TB", target_dtc_price="A$25–40",
         landed_cost_note="Có báo giá Newsun (08/2026) — chưa nhập DB"),
    dict(product_id=ids["no-pull-dog-harness"], assessed_on=D25, weight_class="nhẹ", hazmat="Không",
         policy_risk="TB", retail_risk="TB", bundle_potential="TB", target_dtc_price="A$40–60"),
    dict(product_id=ids["tech-pouch"], assessed_on=D25, weight_class="nhẹ", hazmat="Không", policy_risk="Thấp",
         retail_risk="TB", bundle_potential="Cao", target_dtc_price="A$35–49"),
    dict(product_id=ids["rfid-wallet"], assessed_on=D25, weight_class="nhẹ", hazmat="Không", policy_risk="Thấp",
         retail_risk="TB", bundle_potential="TB", target_dtc_price="A$39–59")])

# ---------- criteria ----------
put("criteria_versions", [
    dict(version="v1", created_on=D23, is_current=False,
         summary="Ngưỡng tô đỏ/xanh của sheet 23/09. Trọng số bằng nhau, không có lọc cứng. Dùng tỷ lệ active/tổng của page (cột AJ)."),
    dict(version="v2", created_on=D25, is_current=True,
         summary="Thêm lọc cứng (pin/lỏng/dao, mẹ&bé, thú cưng, claim, hàng nặng, AOV). Ưu tiên nhu cầu AU (thang log) và 'ads win' đo bằng số ad đang chạy tuyệt đối + số page DTC mạnh. Bỏ tỷ lệ active/tổng và lượt mua tuyệt đối.")])
LVL = "case {c} when 'Thấp' then 1 when 'TB' then 0.5 when 'Cao' then 0 end"
LVL_UP = "case {c} when 'Cao' then 1 when 'TB' then 0.5 when 'Thấp' then 0 end"
W = "case weight_class when 'nhẹ' then 1 when 'TB' then 0.5 when 'nặng' then 0 end"
v1 = [  # key, name, group, expr, good, bad
    ("us_searches", "Search Amazon US", "Nhu cầu", "us_searches", 100000, 10000),
    ("us_purchases", "Lượt mua từ keyword US", "Nhu cầu", "us_purchases", 5000, 500),
    ("us_purchase_rate", "Tỷ lệ mua/search US", "Nhu cầu", "us_purchase_rate", 0.05, 0.015),
    ("au_searches", "Search Amazon AU", "Nhu cầu", "au_searches", 5000, 1000),
    ("uk_searches", "Search Amazon UK", "Nhu cầu", "uk_searches", 20000, 3000),
    ("us_price", "Giá TB US ($)", "Giá & AOV", "us_price", 25, 12),
    ("au_price", "Giá TB AU (A$)", "Giá & AOV", "au_price", 40, 18),
    ("uk_price", "Giá TB UK (£)", "Giá & AOV", "uk_price", 20, 9),
    ("us_click_conc", "Độ tập trung click US", "Cạnh tranh", "us_click_conc", 0.25, 0.45),
    ("tt_units", "TikTok US: SKU top bán 30 ngày", "Xã hội", "tt_units_30d", 5000, 500),
    ("tt_price", "TikTok US: giá SKU top", "Giá & AOV", "tt_price", 30, 10),
    ("tt_views", "TikTok: view video nổi bật", "Xã hội", "tt_views", 100000, 1000),
    ("trend_floor", "Trends: sàn P10/median", "Xu hướng", "trend_floor", 0.8, 0.6),
    ("trend_season", "Trends: mùa vụ max/median", "Xu hướng", "trend_season", 1.6, 3),
    ("trend_growth", "Trends: tăng trưởng 2025/2022", "Xu hướng", "trend_growth", 1.2, 0.95),
    ("meta_active_band", "Meta AU: ad active trong khoảng 100–2000", "Ads win",
     "case when meta_active_ads between 100 and 2000 then 1 when meta_active_ads < 20 or meta_active_ads > 10000 then 0 when meta_active_ads is not null then 0.5 end", 1, 0),
    ("meta_advertisers", "Meta AU: số advertiser (30 ad đầu)", "Ads win", "meta_advertisers", 10, 3),
    ("meta_over60", "Meta AU: ad chạy >60 ngày (30 ad đầu)", "Ads win", "meta_over60", 15, 5),
    ("weight", "Cân nặng", "Vận hành", W, 1, 0),
    ("hazmat", "Không pin/lỏng/dễ vỡ/dao", "Vận hành", "case when hazmat = 'Không' then 1 when hazmat is not null then 0 end", 1, 0),
    ("policy", "Rủi ro chính sách/claim", "Rủi ro", LVL.format(c="policy_risk"), 1, 0),
    ("retail", "Rủi ro siêu thị (Kmart/BIG W)", "Rủi ro", LVL.format(c="retail_risk"), 1, 0),
    ("bundle", "Khả năng bundle/AOV", "Giá & AOV", LVL_UP.format(c="bundle_potential"), 1, 0),
]
crit = [dict(version="v1", key=k, name_vi=n, kind="score", group_name=g, expr=e, good=gd, bad=bd, weight=1,
             status="active", sort=i, rationale="Ngưỡng tô màu của sheet 23/09") for i, (k, n, g, e, gd, bd) in enumerate(v1)]
crit.append(dict(version="v1", key="meta_page_ratio", name_vi="Tỷ lệ ad active/tổng của page (màu cột AJ)",
                 kind="score", group_name="Ads win", expr="null", good=0.8, bad=0.2, weight=0, status="retired", sort=99,
                 rationale="Chỉ dùng để tô màu, không vào điểm. Sai với brand chi nhiều (Guard Blinds 19% nhưng thắng rõ nhất)."))
v2 = [  # key, name, kind, group, expr, good, bad, weight, rationale
    ("hf_hazmat", "Không có pin/chất lỏng/dễ vỡ/lưỡi dao/vũ khí", "hard_filter", "Lọc cứng",
     "coalesce(hazmat, 'Không') = 'Không'", None, None, 0, "Ship quốc tế + chính sách quảng cáo Meta"),
    ("hf_category", "Không thuộc mẹ & bé / thú cưng (giai đoạn đầu)", "hard_filter", "Lọc cứng",
     "coalesce(category, '') not in ('Baby', 'Pet')", None, None, 0,
     "Anh Thanh: mẹ&bé rủi ro an toàn + khách kỹ tính; thú cưng CPM cao, khó cho người mới"),
    ("hf_policy", "Không bán bằng claim y tế/rủi ro chính sách cao", "hard_filter", "Lọc cứng",
     "coalesce(policy_risk, 'Thấp') <> 'Cao'", None, None, 0, "Tránh chết tài khoản quảng cáo khi vốn ít"),
    ("hf_weight", "Không nặng/cồng kềnh", "hard_filter", "Lọc cứng", "coalesce(weight_class, 'TB') <> 'nặng'",
     None, None, 0, "Phí ship AU + đổi trả"),
    ("hf_aov", "Đạt được đơn A$50–100 (giá ≥ A$20 hoặc bundle tốt)", "hard_filter", "Lọc cứng",
     "coalesce(au_price, us_price * 1.55, 0) >= 20 or bundle_potential = 'Cao'", None, None, 0,
     "Giá quá thấp thì CAC ăn hết biên lợi nhuận (1 USD ≈ 1.55 AUD)"),
    ("au_demand", "Nhu cầu AU (search Amazon AU, thang log)", "score", "Nhu cầu", "ln(nullif(au_searches, 0))",
     9.9, 6.9, 3, "Thị trường chính là AU; 1k→0 điểm, 20k→100 điểm"),
    ("us_demand", "Nhu cầu US (thang log)", "score", "Nhu cầu", "ln(nullif(us_searches, 0))", 12.6, 9.9, 1,
     "Tham chiếu quy mô; 20k→0, 300k→100"),
    ("uk_demand", "Nhu cầu UK (thang log)", "score", "Nhu cầu", "ln(nullif(uk_searches, 0))", 10.8, 8.0, 1,
     "Thị trường mở rộng; 3k→0, 50k→100"),
    ("us_purchase_rate", "Tỷ lệ mua/search US", "score", "Nhu cầu", "us_purchase_rate", 0.04, 0.01, 1.5,
     "Keyword có ý định mua; không dùng lượt mua tuyệt đối vì chỉ là lượt mua từ đúng 1 keyword"),
    ("au_price", "Giá TB AU (A$)", "score", "Giá & AOV", "au_price", 40, 18, 2, "Đủ biên cho CAC Meta"),
    ("bundle", "Khả năng bundle/AOV", "score", "Giá & AOV", LVL_UP.format(c="bundle_potential"), 1, 0, 1.5,
     "Bán theo bộ để đạt AOV"),
    ("us_click_conc", "Độ tập trung click US (thấp = phân tán)", "score", "Cạnh tranh", "us_click_conc", 0.25, 0.45, 1,
     "Không bị vài SKU chiếm"),
    ("marketplace", "Không bị Temu/siêu thị chạy ad chiếm", "score", "Cạnh tranh", "marketplace_pages", 0, 2, 1,
     "Số page Temu/BIG W/Petbarn... trong nhóm advertiser"),
    ("retail", "Rủi ro siêu thị (Kmart/BIG W)", "score", "Cạnh tranh", LVL.format(c="retail_risk"), 1, 0, 1.5,
     "Hàng có sẵn giá rẻ ở siêu thị AU"),
    ("adswin_pages", "Số page DTC mạnh (≥50 ad đang chạy)", "score", "Ads win", "adswin_pages", 3, 0, 3,
     "Tiêu chí anh Thanh; đo bằng số ad tuyệt đối thay vì tỷ lệ active/tổng"),
    ("meta_over60", "Ad chạy >60 ngày (trong ~30 ad đầu, AU)", "score", "Ads win", "meta_over60", 15, 5, 2,
     "Ad chạy lâu = có lãi"),
    ("weight", "Cân nặng", "score", "Vận hành", W, 1, 0, 1, "Phí ship"),
    ("trend_floor", "Trends: nhu cầu nền ổn định", "score", "Xu hướng", "trend_floor", 0.8, 0.6, 1, "Ít mùa vụ"),
    ("trend_season", "Trends: mùa vụ (thấp = tốt)", "score", "Xu hướng", "trend_season", 1.6, 3, 1, None),
    ("trend_growth", "Trends: tăng trưởng 2025/2022", "score", "Xu hướng", "trend_growth", 1.2, 0.95, 0.5, None),
    ("tt_units", "TikTok US: SKU top bán 30 ngày", "score", "Xã hội", "tt_units_30d", 5000, 500, 0.5,
     "Tín hiệu phụ (chỉ có US)"),
]
crit += [dict(version="v2", key=k, name_vi=n, kind=kd, group_name=g, expr=e, good=gd, bad=bd, weight=w,
              status="active", sort=i, rationale=ra) for i, (k, n, kd, g, e, gd, bd, w, ra) in enumerate(v2)]
crit += [
    dict(version="v2", key="us_purchases", name_vi="Lượt mua từ keyword US (tuyệt đối)", kind="score", group_name="Nhu cầu",
         expr="us_purchases", good=5000, bad=500, weight=0, status="retired", sort=90,
         rationale="Bỏ 25/09: con số chỉ là lượt mua từ đúng 1 keyword (packing cubes 14,936 vs tổng thị trường ~250k/tháng)"),
    dict(version="v2", key="meta_active_band", name_vi="Meta AU: ad active trong khoảng 100–2000", kind="score",
         group_name="Ads win", expr="null", good=1, bad=0, weight=0, status="retired", sort=91,
         rationale="Thay bằng adswin_pages + meta_over60"),
    dict(version="v2", key="meta_page_ratio", name_vi="Tỷ lệ ad active/tổng của page", kind="score", group_name="Ads win",
         expr="null", good=0.8, bad=0.2, weight=0, status="retired", sort=92,
         rationale="Sai với brand test nhiều creative (Guard Blinds 501/2642 = 19% nhưng là brand thắng)"),
    dict(version="v2", key="us_price", name_vi="Giá TB US", kind="score", group_name="Giá & AOV", expr="us_price",
         good=25, bad=12, weight=0, status="retired", sort=93, rationale="Thị trường chính là AU → dùng au_price"),
]
put("criteria", crit)

# ---------- decisions + dossiers ----------
def to_markdown(body):
    out = []
    for l in body.split("\n"):
        if l.startswith("KẾT LUẬN:"):
            l = "**KẾT LUẬN:**" + l[len("KẾT LUẬN:"):]
        elif l.startswith("• "):
            l = "- " + l[2:]
        elif l.startswith("− "):
            l = "- **Rủi ro:** " + l[2:].removeprefix("Rủi ro: ").removeprefix("Rủi ro: ")
        elif l.startswith("→ "):
            l = "\n" + l
        out.append(l)
    return "\n".join(out)


dec_rows, dos_rows = [], []
for r in pipe:
    no = int(r["#"])
    d, body = E[no]
    pid = by_no[no]
    head = body.split("\n")[0].replace("KẾT LUẬN: ", "")
    dec_rows.append(dict(product_id=pid, decided_on=D25, decision=DECISION[d], reason=head, criteria_version="v2",
                         revisit_trigger="Có báo giá landed cost AU / số ad Meta thay đổi mạnh" if DECISION[d] != "khong_chon" else None))
    dos_rows.append(dict(product_id=pid, version=1, written_on=D25, verdict=d, headline=head,
                         body_md=to_markdown(body)))
DOSSIER_NEW = json.loads((Path(__file__).parent / "data" / "dossiers_2026_09_25.json").read_text())
for slug, x in DOSSIER_NEW.items():
    dos_rows.append(dict(product_id=ids[slug], version=1, written_on=D25, verdict=x["verdict"], headline=x["headline"],
                         body_md=x["body_md"]))
    if x.get("decision"):
        dec_rows.append(dict(product_id=ids[slug], decided_on=D25, decision=x["decision"], reason=x["headline"],
                             criteria_version="v2", revisit_trigger=x.get("revisit")))
put("decisions", dec_rows)
put("dossiers", dos_rows)

put("raw_files", [
    dict(captured_on=D23, source="topview + meta (pipeline 91 SP)", path=str(RAW.relative_to(ROOT.parent)), note="JSON/TSV gốc 23/09"),
    dict(captured_on=D23, source="xlsx", path=str(XLSX), note="Pipeline 91 SP + sheet dữ liệu thô"),
    dict(captured_on=D25, source="topview + meta + amazon", path=str((RAW / "2026-09-25.json").relative_to(ROOT.parent)),
         note="EDC, teething, harness, advertiser snapshots, listing BAGAIL")])
print("done")
